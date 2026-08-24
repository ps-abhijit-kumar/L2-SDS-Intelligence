import os
import sys

# Ensure project root is on sys.path for subprocess invocations
project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if project_root not in sys.path:
    sys.path.insert(0, project_root)

import json
import openpyxl
import pandas as pd
from typing import Optional, Dict, Any, List, Set
from mcp.server.fastmcp import FastMCP

from src.workbook_utils import (
    COLUMN_SYNONYMS,
    INVALID_PRODUCT_VALUES,
    is_valid_product_value,
    detect_column_mapping,
    classify_sheet,
    inspect_workbook
)

mcp = FastMCP("ExcelMCP")

# Target file configuration
EXCEL_FILE = os.getenv("EXCEL_FILE_PATH", "sample_requests_eval.xlsx")

@mcp.tool()
def get_pending_requests(column_mapping_json: str = "", selected_sheets_json: str = "") -> str:
    """Reads all semantically normalized pending SDS request records from the active workbook."""
    target_path = os.getenv("EXCEL_FILE_PATH", EXCEL_FILE)
    if not os.path.exists(target_path):
        return "[]"

    custom_mapping = None
    if column_mapping_json:
        try:
            custom_mapping = json.loads(column_mapping_json)
        except Exception:
            pass

    custom_selected_sheets = None
    if selected_sheets_json:
        try:
            parsed_sheets = json.loads(selected_sheets_json)
            if isinstance(parsed_sheets, list) and len(parsed_sheets) > 0:
                custom_selected_sheets = parsed_sheets
        except Exception:
            pass

    try:
        preview = inspect_workbook(
            target_path,
            custom_mapping=custom_mapping,
            custom_selected_sheets=custom_selected_sheets
        )
        all_rows = preview.get("rows", [])
        pending = [r for r in all_rows if str(r.get("Status", "")).upper() != "EXACT MATCH"]
        return json.dumps(pending)
    except Exception as e:
        return "[]"

@mcp.tool()
def update_request_status(
    row_index: int = 0,
    final_url: str = "",
    status: str = "",
    confidence: int = 0,
    reasoning: str = "",
    sheet_name: str = "",
    excel_row: int = 0
) -> str:
    """Updates the specified row in the specified Excel sheet with the results."""
    target_path = os.getenv("EXCEL_FILE_PATH", EXCEL_FILE)
    if not os.path.exists(target_path):
        return "Error: File not found."

    wb = openpyxl.load_workbook(target_path)

    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.active

    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[1]) if cell.value}

    for col_name in ["Found URL", "Status", "Confidence", "Reasoning"]:
        if col_name not in headers:
            col_num = ws.max_column + 1
            ws.cell(row=1, column=col_num, value=col_name)
            headers[col_name] = col_num

    if excel_row > 0:
        target_row = excel_row
    else:
        target_row = row_index + 2

    ws.cell(row=target_row, column=headers["Found URL"], value=final_url)
    ws.cell(row=target_row, column=headers["Status"], value=status)
    ws.cell(row=target_row, column=headers["Confidence"], value=confidence)
    ws.cell(row=target_row, column=headers["Reasoning"], value=reasoning)

    wb.save(target_path)
    return "Success"

if __name__ == "__main__":
    mcp.run()
