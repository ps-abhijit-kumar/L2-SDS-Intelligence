import pytest
import io
import os
import openpyxl
import pandas as pd
from fastapi.testclient import TestClient
from server import app, inspect_workbook, detect_column_mapping, classify_sheet, DEFAULT_EXCEL_FILE

@pytest.fixture
def client():
    return TestClient(app)

def test_health_endpoint(client):
    res = client.get('/api/health')
    assert res.status_code == 200
    data = res.json()
    assert data['status'] == 'healthy'
    assert data['backend'] == 'operational'
    assert 'mcp_server_available' in data
    assert 'batch_status' in data

def test_batch_preview_default(client):
    res = client.get('/api/batch/preview')
    assert res.status_code == 200
    data = res.json()
    assert data['file_name'] is None
    assert data['sds_requests_count'] == 0
    assert data['sds_sheets_count'] == 0
    assert len(data['rows']) == 0

def test_sheet_classification():
    # SDS Requests sheet
    assert classify_sheet('Part1', ['Product Product Name', 'Product Company Name', 'Country'], 20, 20) == 'SDS_REQUESTS'
    assert classify_sheet('Part10', ['Product Name', 'Manufacturer'], 15, 15) == 'SDS_REQUESTS'
    assert classify_sheet('Sheet1', ['Product', 'Company'], 10, 10) == 'SDS_REQUESTS'

    # Supporting data sheet
    assert classify_sheet('Allocation', ['Squad', 'Member', 'Count'], 0, 897) == 'SUPPORTING_DATA'
    assert classify_sheet('Data_Dump', ['Ref_ID', 'Val1', 'Val2'], 0, 5336) == 'SUPPORTING_DATA'

    # Summary sheet
    assert classify_sheet('Summary_Report', ['Metric', 'Total'], 0, 15) == 'SUMMARY'
    assert classify_sheet('Pivot_Table', ['Category', 'Sum'], 0, 30) == 'SUMMARY'

def test_multi_sheet_with_supporting_sheets_filtering(client):
    """
    Verifies multi-dataset workbook upload and single request-set selection:
    1. Upload workbook with multiple candidate SDS request datasets (Part1, Part2, Part3) + supporting + summary sheets.
    2. Dynamic inspection detects all 3 valid request datasets without merging them into thousands of active requests.
    3. User selects ONE request dataset (Part1 -> 20 requests).
    4. Only the selected dataset becomes active (Active=20, Pending=20).
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # remove default

    # 1. Add 3 Part request sheets (20 products each = 60 potential requests across 3 distinct sets)
    for sname in ['Part1', 'Part2', 'Part3']:
        ws = wb.create_sheet(title=sname)
        headers = [
            'S.No.', 'Obtainment Request Id', 'Product Product Name', 'Product Company Name',
            'Part Numbers', 'Language', 'Country'
        ]
        ws.append(headers)
        for i in range(1, 21):
            ws.append([
                i, f'REQ_{sname}_{i}', f'{sname} Chemical Compound {i}',
                f'{sname} Chemical Corp', f'PN-00{i}', 'English', 'United States'
            ])

    # 2. Add Supporting Data sheet (100 rows of non-chemical IDs)
    ws_alloc = wb.create_sheet(title='Allocation')
    ws_alloc.append(['Squad ID', 'Agent Name', 'Assigned Work Items', 'Status'])
    for i in range(1, 101):
        ws_alloc.append([f'SQ_{i}', f'Agent_{i}', i * 5, 'Active'])

    # 3. Add Summary sheet (10 rows)
    ws_sum = wb.create_sheet(title='Summary_Totals')
    ws_sum.append(['Metric Name', 'Total Value'])
    for i in range(1, 11):
        ws_sum.append([f'Metric_{i}', i * 100])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    files = {
        'file': ('Obtainment_NWL_Squads_1(3).xlsx', buffer, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    }

    res = client.post('/api/batch/upload', files=files)
    assert res.status_code == 200
    data = res.json()

    # Total workbook physical rows = 60 + 100 + 10 = 170 across 5 sheets
    assert data['workbook_sheets_count'] == 5
    assert data['workbook_physical_rows'] == 170

    # Inspection detects all 3 valid SDS Request candidate sheets in sheets_summary
    sds_sheets_detected = [s for s in data['sheets_summary'] if s['classification'] == 'SDS_REQUESTS']
    assert len(sds_sheets_detected) == 3
    for s in sds_sheets_detected:
        assert s['valid_requests'] == 20

    # When multiple request datasets exist, the system requires user selection and does NOT auto-merge them
    assert data['sds_sheets_count'] == 0
    assert data['sds_requests_count'] == 0
    assert len(data['rows']) == 0

    # Verify column mapping auto-detected
    assert data['column_mapping']['product'] == 'Product Product Name'
    assert data['column_mapping']['company'] == 'Product Company Name'
    assert data['column_mapping']['part_number'] == 'Part Numbers'

    # User selects ONE request dataset (Part1 -> 20 requests)
    confirm_res = client.post('/api/batch/confirm-mapping', json={
        'product': 'Product Product Name',
        'company': 'Product Company Name',
        'part_number': 'Part Numbers',
        'language': 'Language',
        'country': 'Country',
        'selected_sheets': ['Part1']
    })
    assert confirm_res.status_code == 200
    confirm_data = confirm_res.json()

    # Only the selected dataset (Part1) becomes active
    assert confirm_data['sds_sheets_count'] == 1
    assert confirm_data['sds_requests_count'] == 20
    assert confirm_data['pending_requests_count'] == 20
    assert len(confirm_data['rows']) == 20
    assert all(r['_sheet_name'] == 'Part1' for r in confirm_data['rows'])

    # Restore default benchmark
    res_default = client.post('/api/batch/select-default')
    assert res_default.status_code == 200
    assert res_default.json()['sds_requests_count'] == 10

def test_batch_export(client):
    res = client.get('/api/batch/export')
    assert res.status_code == 200
    assert 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in res.headers.get('content-type', '')
