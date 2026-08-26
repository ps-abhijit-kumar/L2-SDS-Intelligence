import pytest
import os
import json
import tempfile
import openpyxl
from src.mcp_client import SDSMCPClient

@pytest.mark.asyncio
async def test_mcp_client_isolated_roundtrip():
    # Create isolated temporary Excel workbook (Priority 8: never mutate committed benchmark files)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        # Build sample test workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["S.No.", "Product Name", "Product Company Name", "Part Number", "Language", "Country"])
        ws.append([1, "Acetone", "Sigma-Aldrich", "179124", "English", "United States"])
        wb.save(tmp_path)

        client = SDSMCPClient(excel_file_path=tmp_path)
        await client.connect()
        try:
            # 1. Get pending requests
            pending_json = await client.get_pending_requests()
            requests = json.loads(pending_json)
            assert isinstance(requests, list)
            assert len(requests) == 1
            assert requests[0]["Product Name"] == "Acetone"

            # 2. Update request status via stdio MCP tool
            row_idx = requests[0].get("_row_index", 0)
            res = await client.update_request_status(
                row_index=row_idx,
                final_url="https://www.sigmaaldrich.com/sds/179124.pdf",
                status="EXACT MATCH",
                confidence=95,
                reasoning="Isolated test verified result",
                sheet_name="Sheet1",
                excel_row=2
            )
            assert "Success" in res

            # Verify in-place update in isolated workbook
            wb_updated = openpyxl.load_workbook(tmp_path)
            ws_up = wb_updated["Sheet1"]
            headers = [cell.value for cell in ws_up[1]]
            assert "Found URL" in headers
            assert "Status" in headers
            assert "Confidence" in headers
            assert "Reasoning" in headers

            # Check row 2 values
            status_col = headers.index("Status") + 1
            assert ws_up.cell(row=2, column=status_col).value == "EXACT MATCH"
        finally:
            await client.disconnect()
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)

@pytest.mark.asyncio
async def test_mcp_inspect_sds_document_isolated():
    client = SDSMCPClient()
    await client.connect()
    try:
        # Inspect an invalid URL to test error handling & SSRF defense over stdio MCP
        res_json = await client.inspect_sds_document("http://127.0.0.1:9999/malicious.pdf")
        data = json.loads(res_json)
        assert isinstance(data, dict)
        assert data.get("fetched_successfully") is False
        assert "Security Error" in data.get("error", "") or "error" in data
    finally:
        await client.disconnect()
